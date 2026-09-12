#!/usr/bin/env python3
"""
Script to create a PowerPoint presentation with the Child Law Project architecture diagram
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import os

def create_architecture_ppt():
    """Create a PowerPoint presentation with the architecture diagram"""
    
    # Create presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define color scheme
    COLOR_BLUE = RGBColor(225, 245, 255)  # Light blue
    COLOR_ORANGE = RGBColor(255, 243, 224)  # Light orange
    COLOR_PURPLE = RGBColor(243, 229, 245)  # Light purple
    COLOR_RED = RGBColor(255, 235, 238)  # Light red
    COLOR_TEXT = RGBColor(33, 33, 33)  # Dark text
    COLOR_ACCENT = RGBColor(13, 110, 253)  # Blue accent
    
    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide1 = prs.slides.add_slide(title_slide_layout)
    background = slide1.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(13, 71, 161)  # Dark blue
    
    # Add title
    title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = "Child Law Legal Assistant"
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(9), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    p = subtitle_frame.paragraphs[0]
    p.text = "System Architecture & Workflow"
    p.font.size = Pt(32)
    p.font.color.rgb = RGBColor(200, 220, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Architecture Overview
    blank_slide_layout = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(blank_slide_layout)
    background = slide2.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)  # Light gray
    
    # Add title
    title_box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "System Architecture Overview"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(13, 110, 253)
    
    # Add architecture components
    components = [
        {
            "title": "Input Layer",
            "items": ["app.py - Main Entry Point", "Interactive Menu", "User Questions"],
            "color": COLOR_PURPLE,
            "x": 0.5,
            "y": 1.2
        },
        {
            "title": "Validation Layer",
            "items": ["answer_generation.py", "Domain Validation", "State Filtering"],
            "color": COLOR_ORANGE,
            "x": 3.5,
            "y": 1.2
        },
        {
            "title": "Data Processing",
            "items": ["document_processing.py", "Load & Chunk Documents", "embeddings.py - Cache Management"],
            "color": COLOR_BLUE,
            "x": 6.5,
            "y": 1.2
        },
        {
            "title": "Retrieval Methods",
            "items": ["retrieval.py", "Dense, Hybrid,", "Multi-Query, Corrective"],
            "color": COLOR_ORANGE,
            "x": 0.5,
            "y": 3.5
        },
        {
            "title": "Generation & Processing",
            "items": ["answer_generation.py", "LLM-based Answer Gen", "Citation Management"],
            "color": COLOR_BLUE,
            "x": 3.5,
            "y": 3.5
        },
        {
            "title": "Evaluation",
            "items": ["evaluation.py", "Relevance Scoring", "Performance Metrics"],
            "color": COLOR_BLUE,
            "x": 6.5,
            "y": 3.5
        },
    ]
    
    for comp in components:
        # Add rectangle shape
        shape = slide2.shapes.add_shape(
            1,  # Rectangle
            Inches(comp["x"]),
            Inches(comp["y"]),
            Inches(2.7),
            Inches(1.6)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = comp["color"]
        shape.line.color.rgb = RGBColor(150, 150, 150)
        shape.line.width = Pt(2)
        
        # Add title text
        text_frame = shape.text_frame
        text_frame.word_wrap = True
        text_frame.margin_top = Inches(0.1)
        text_frame.margin_left = Inches(0.1)
        text_frame.margin_right = Inches(0.1)
        
        p = text_frame.paragraphs[0]
        p.text = comp["title"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT
        p.alignment = PP_ALIGN.CENTER
        
        # Add items
        for item in comp["items"]:
            p = text_frame.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT
            p.space_before = Pt(2)
            p.level = 0
    
    # Slide 3: Data Flow
    slide3 = prs.slides.add_slide(blank_slide_layout)
    background = slide3.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)
    
    # Add title
    title_box = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Data Flow & Processing Pipeline"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(13, 110, 253)
    
    # Add flow steps
    flow_steps = [
        ("1. Question Input", "User enters a question about Child Law"),
        ("2. Validation", "Check if question is related to Child Laws & filter by state"),
        ("3. Embeddings", "Load or create embeddings cache for documents"),
        ("4. Retrieval", "Use Dense/Hybrid/Multi-Query/Corrective methods to retrieve relevant chunks"),
        ("5. Generation", "Generate answer using llama3.2 LLM with citations"),
        ("6. Evaluation", "Score relevance and calculate performance metrics"),
        ("7. Display", "Show answer with citations and source documents"),
    ]
    
    flow_box = slide3.shapes.add_textbox(Inches(1), Inches(1.2), Inches(8), Inches(5.5))
    text_frame = flow_box.text_frame
    text_frame.word_wrap = True
    
    for i, (step, description) in enumerate(flow_steps):
        if i > 0:
            p = text_frame.add_paragraph()
            p.text = ""
            p.space_after = Pt(3)
        
        p = text_frame.add_paragraph()
        p.text = step
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.space_after = Pt(2)
        
        p = text_frame.add_paragraph()
        p.text = description
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.level = 1
        p.space_after = Pt(6)
    
    # Slide 4: Retrieval Methods
    slide4 = prs.slides.add_slide(blank_slide_layout)
    background = slide4.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)
    
    # Add title
    title_box = slide4.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Retrieval Methods"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(13, 110, 253)
    
    retrieval_methods = [
        ("Dense Retrieval", "Uses embedding similarity search to find top-K relevant chunks"),
        ("Hybrid Retrieval", "Combines 70% embedding similarity + 30% BM25 keyword matching"),
        ("Multi-Query Retrieval", "Generates multiple query variants and aggregates results"),
        ("Corrective Retrieval", "Fallback method with quality validation for better accuracy"),
    ]
    
    y_position = 1.3
    for method, description in retrieval_methods:
        # Add method name
        method_box = slide4.shapes.add_textbox(Inches(1), Inches(y_position), Inches(8), Inches(0.3))
        text_frame = method_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = "▸ " + method
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        
        # Add description
        desc_box = slide4.shapes.add_textbox(Inches(1.3), Inches(y_position + 0.35), Inches(7.7), Inches(0.8))
        text_frame = desc_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = description
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        
        y_position += 1.3
    
    # Slide 5: Technical Stack
    slide5 = prs.slides.add_slide(blank_slide_layout)
    background = slide5.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(245, 245, 245)
    
    # Add title
    title_box = slide5.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = "Technical Stack & Models"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(13, 110, 253)
    
    # Add technical details in columns
    tech_left = [
        ("Embedding Model", "nomic-embed-text"),
        ("Language Model", "llama3.2"),
        ("Framework", "Python with LangChain"),
        ("Vector Storage", "In-memory embeddings cache"),
    ]
    
    tech_right = [
        ("Document Format", "PDF & Text files"),
        ("Keyword Matching", "BM25 (rank_bm25)"),
        ("API Integration", "Ollama (local LLM)"),
        ("Main Dependencies", "OpenAI SDK, LangChain"),
    ]
    
    # Left column
    y = 1.3
    for label, value in tech_left:
        label_box = slide5.shapes.add_textbox(Inches(0.7), Inches(y), Inches(2), Inches(0.3))
        text_frame = label_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = label + ":"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        
        value_box = slide5.shapes.add_textbox(Inches(0.7), Inches(y + 0.3), Inches(2), Inches(0.4))
        text_frame = value_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = value
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        
        y += 0.9
    
    # Right column
    y = 1.3
    for label, value in tech_right:
        label_box = slide5.shapes.add_textbox(Inches(5.3), Inches(y), Inches(2.2), Inches(0.3))
        text_frame = label_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = label + ":"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        
        value_box = slide5.shapes.add_textbox(Inches(5.3), Inches(y + 0.3), Inches(3.8), Inches(0.4))
        text_frame = value_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = value
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        
        y += 0.9
    
    # Save presentation
    output_path = os.path.join(
        os.path.dirname(__file__),
        "Child_Law_Architecture.pptx"
    )
    
    prs.save(output_path)
    print(f"[OK] PowerPoint presentation created successfully!")
    print(f"[OK] Saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    create_architecture_ppt()
